import os
import re
import math
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from collections import Counter
from datetime import datetime, date, timedelta

from openpyxl import load_workbook
from openpyxl.styles import PatternFill


# ============================================================
# COLORS
# ============================================================

YELLOW_FILL = PatternFill(
    fill_type="solid",
    fgColor="FFFF00"
)

BLUE_FILL = PatternFill(
    fill_type="solid",
    fgColor="0000FF"
)

PINK_FILL = PatternFill(
    fill_type="solid",
    fgColor="FFC0CB"
)


# ============================================================
# TASKS
# ============================================================

TASKS = [
    "Trim",
    "Poper Case",
    "Upper Case",
    "Remove Special Character",
    "SSN Format",
    "ITIN Find",
    "Find Invalid SSN",
    "DOB (mm/dd/yyyy)",
    "Phone Number (123-456-7890)",
    "Street Address Format",
    "Length Count",
]


# ============================================================
# REMOVE SPECIAL CHARACTER TASK
# ============================================================

REMOVE_SPECIAL_CHARS = set(
    list('!@#$%^&_+=[]{}:;"./\\?<>`~')
)


# ============================================================
# STREET ADDRESS ABBREVIATIONS
# ============================================================

STREET_ABBREVIATIONS = {
    "Street": "St",
    "Road": "Rd",
    "Avenue": "Ave",
    "Boulevard": "Blvd",
    "Drive": "Dr",
    "Lane": "Ln",
    "Court": "Ct",
    "Circle": "Cir",
    "Place": "Pl",
    "Parkway": "Pkwy",
    "North": "N",
    "South": "S",
    "East": "E",
    "West": "W",
    "Northeast": "NE",
    "Northwest": "NW",
    "Southeast": "SE",
    "Southwest": "SW",
    "Apartment": "Apt",
    "Suite": "Ste",
    "Po Box":"PO Box",
    "P O Box":"PO Box",
    "P o Box":"PO Box",  
}


# ============================================================
# HELPER FUNCTIONS
# ============================================================

def is_blank(value):
    if value is None:
        return True

    if isinstance(value, str):
        return value.strip() == ""

    return False


def value_to_text(value):
    if value is None:
        return ""

    if isinstance(value, datetime):
        return value.strftime("%m/%d/%Y")

    if isinstance(value, date):
        return value.strftime("%m/%d/%Y")

    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))

    return str(value)


def normalize_digits(value):
    if value is None:
        return ""

    return re.sub(r"\D", "", str(value))


def trim_excel_style(value):
    if value is None:
        return value

    if not isinstance(value, str):
        return value

    return re.sub(r"\s+", " ", value).strip()


def proper_case(value):
    if value is None:
        return value

    if not isinstance(value, str):
        return value

    return value.lower().title()


def upper_case(value):
    if value is None:
        return value

    if not isinstance(value, str):
        return value

    return value.upper()


def remove_selected_specials(value):
    if value is None:
        return value

    if not isinstance(value, str):
        value = str(value)

    result = []

    for char in value:
        if char in REMOVE_SPECIAL_CHARS:
            result.append(" ")
        else:
            result.append(char)

    return "".join(result)


# ============================================================
# SSN FORMAT
# ============================================================

def format_ssn(value):

    if is_blank(value):
        return value

    digits = normalize_digits(value)

    if not digits:
        return value

    if len(digits) > 9:
        return value

    digits = digits.zfill(9)

    return (
        digits[:3]
        + "-"
        + digits[3:5]
        + "-"
        + digits[5:]
    )


# ============================================================
# PHONE FORMAT
# ============================================================

def format_phone(value):

    if is_blank(value):
        return value

    digits = normalize_digits(value)

    if not digits:
        return value

    if len(digits) > 10:
        return value

    digits = digits.zfill(10)

    return (
        digits[:3]
        + "-"
        + digits[3:6]
        + "-"
        + digits[6:]
    )


# ============================================================
# DATE PARSER
# ============================================================

def parse_date_value(value):

    if value is None:
        return None

    if isinstance(value, datetime):
        return value

    if isinstance(value, date):
        return datetime(
            value.year,
            value.month,
            value.day
        )

    # Excel serial number
    if isinstance(value, (int, float)):

        if not isinstance(value, bool):

            try:

                number = float(value)

                if 1 <= number <= 2958465:

                    base = datetime(
                        1899,
                        12,
                        30
                    )

                    return base + timedelta(
                        days=number
                    )

            except Exception:
                pass

    if not isinstance(value, str):
        return None

    text = value.strip()

    if not text:
        return None

    # Numeric string Excel serial
    if re.fullmatch(
        r"\d+(\.\d+)?",
        text
    ):

        try:

            number = float(text)

            if 1 <= number <= 2958465:

                base = datetime(
                    1899,
                    12,
                    30
                )

                return base + timedelta(
                    days=number
                )

        except Exception:
            pass

    formats = [
        "%m/%d/%Y",
        "%m-%d-%Y",
        "%m.%d.%Y",

        "%m/%d/%y",
        "%m-%d-%y",
        "%m.%d.%y",

        "%Y/%m/%d",
        "%Y-%m-%d",
        "%Y.%m.%d",

        "%d/%m/%Y",
        "%d-%m-%Y",
        "%d.%m.%Y",

        "%d/%m/%y",
        "%d-%m-%y",
        "%d.%m.%y",

        "%B %d, %Y",
        "%b %d, %Y",
        "%B %d %Y",
        "%b %d %Y",
    ]

    for fmt in formats:

        try:
            return datetime.strptime(
                text,
                fmt
            )

        except ValueError:
            continue

    return None


# ============================================================
# DOB FORMAT
# ============================================================

def format_dob(value):

    if is_blank(value):
        return value

    parsed = parse_date_value(value)

    if parsed is None:
        return value

    return parsed.strftime(
        "%m/%d/%Y"
    )


# ============================================================
# STREET ADDRESS FORMAT
# ============================================================

def street_address_format(value):

    if is_blank(value):
        return value

    text = value_to_text(value)

    # --------------------------------------------------------
    # Replace ALL special characters with spaces.
    # --------------------------------------------------------

    text = re.sub(
        r"[^A-Za-z0-9\s]",
        " ",
        text
    )

    # --------------------------------------------------------
    # Collapse spaces
    # --------------------------------------------------------

    text = re.sub(
        r"\s+",
        " ",
        text
    ).strip()

    if not text:
        return text

    # --------------------------------------------------------
    # Proper Case
    # --------------------------------------------------------

    text = text.lower().title()

    # --------------------------------------------------------
    # Ordinal normalization
    #
    # 45Th  -> 45th
    # 45TH  -> 45th
    # 45th  -> 45th
    # 21ST  -> 21st
    # 22ND  -> 22nd
    # 23RD  -> 23rd
    # --------------------------------------------------------

    def normalize_ordinal(match):

        number = match.group(1)
        suffix = match.group(2).lower()

        return number + suffix

    text = re.sub(
        r"\b(\d+)(st|nd|rd|th)\b",
        normalize_ordinal,
        text,
        flags=re.IGNORECASE
    )

    # --------------------------------------------------------
    # Address abbreviations
    # --------------------------------------------------------

    for full_word, abbreviation in STREET_ABBREVIATIONS.items():

        pattern = (
            r"\b"
            + re.escape(full_word)
            + r"\b"
        )

        text = re.sub(
            pattern,
            abbreviation,
            text,
            flags=re.IGNORECASE
        )

    # --------------------------------------------------------
    # Final cleanup
    # --------------------------------------------------------

    text = re.sub(
        r"\s+",
        " ",
        text
    ).strip()

    return text


# ============================================================
# ITIN FIND
# ============================================================

def starts_with_nine(value):

    if is_blank(value):
        return False

    text = str(value).strip()

    return text.startswith("9")


# ============================================================
# INVALID SSN
# ============================================================

def invalid_ssn(value):

    if is_blank(value):
        return False

    digits = normalize_digits(value)

    if len(digits) != 9:
        return False

    first_three = digits[:3]
    middle_two = digits[3:5]
    last_four = digits[5:]

    if first_three == "666":
        return True

    if first_three == "000":
        return True

    if middle_two == "00":
        return True

    if last_four == "0000":
        return True

    if digits == "999999999":
        return True

    return False


# ============================================================
# NUMERIC DETECTION
# ============================================================

def is_numeric_value(value):

    if value is None:
        return False

    if isinstance(value, bool):
        return False

    if isinstance(value, (int, float)):

        return not (
            isinstance(value, float)
            and math.isnan(value)
        )

    if isinstance(value, str):

        text = value.strip()

        if not text:
            return False

        return bool(
            re.fullmatch(
                r"[-+]?\d+(\.\d+)?",
                text
            )
        )

    return False


# ============================================================
# MAIN APPLICATION
# ============================================================

class ExcelFormatApp:

    def __init__(self, root):

        self.root = root

        self.root.title(
            "Format Excel File"
        )

        self.root.geometry(
            "1200x750"
        )

        self.root.minsize(
            1050,
            680
        )

        self.input_file = ""
        self.output_file = ""

        self.workbook = None

        self.columns = []

        self.format_results = []

        self.current_column = None

        self.analysis_window = None
        self.format_window = None

        self.build_step1()

    # ========================================================
    # STEP 1
    # ========================================================

    def build_step1(self):

        for widget in self.root.winfo_children():
            widget.destroy()

        title = tk.Label(
            self.root,
            text="Format Excel File",
            font=("Segoe UI", 22, "bold")
        )

        title.pack(
            pady=(30, 5)
        )

        subtitle = tk.Label(
            self.root,
            text="Step 1 - File Setup",
            font=("Segoe UI", 13)
        )

        subtitle.pack(
            pady=(0, 30)
        )

        main = tk.Frame(
            self.root
        )

        main.pack(
            fill="x",
            padx=70
        )

        # ----------------------------------------------------
        # INPUT
        # ----------------------------------------------------

        tk.Label(
            main,
            text="Select Input Excel File:",
            font=("Segoe UI", 11, "bold")
        ).grid(
            row=0,
            column=0,
            sticky="w",
            pady=15
        )

        self.input_var = tk.StringVar()

        ttk.Entry(
            main,
            textvariable=self.input_var,
            width=75
        ).grid(
            row=0,
            column=1,
            padx=10
        )

        ttk.Button(
            main,
            text="Browse File",
            command=self.select_input_file
        ).grid(
            row=0,
            column=2
        )

        # ----------------------------------------------------
        # OUTPUT
        # ----------------------------------------------------

        tk.Label(
            main,
            text="Final Export / Save As:",
            font=("Segoe UI", 11, "bold")
        ).grid(
            row=1,
            column=0,
            sticky="w",
            pady=15
        )

        self.output_var = tk.StringVar()

        ttk.Entry(
            main,
            textvariable=self.output_var,
            width=75
        ).grid(
            row=1,
            column=1,
            padx=10
        )

        ttk.Button(
            main,
            text="Browse File",
            command=self.select_output_file
        ).grid(
            row=1,
            column=2
        )

        ttk.Button(
            self.root,
            text="Get Analysis Data",
            command=self.get_analysis_data
        ).pack(
            pady=45
        )

        tk.Label(
            self.root,
            text=(
                "The original Excel file will never be modified. "
                "Formatting will be exported to the selected file."
            ),
            font=("Segoe UI", 10),
            fg="gray"
        ).pack()

    # ========================================================
    # INPUT FILE
    # ========================================================

    def select_input_file(self):

        path = filedialog.askopenfilename(
            title="Select Input Excel File",
            filetypes=[
                (
                    "Excel Files",
                    "*.xlsx *.xlsm"
                )
            ]
        )

        if path:

            self.input_file = path
            self.input_var.set(path)

            directory = os.path.dirname(path)

            filename = os.path.basename(path)

            name, ext = os.path.splitext(
                filename
            )

            suggested = os.path.join(
                directory,
                name + "_Formatted" + ext
            )

            self.output_file = suggested

            self.output_var.set(
                suggested
            )

    # ========================================================
    # OUTPUT FILE
    # ========================================================

    def select_output_file(self):

        path = filedialog.asksaveasfilename(
            title="Select Final Export File",
            defaultextension=".xlsx",
            filetypes=[
                (
                    "Excel Workbook",
                    "*.xlsx"
                ),
                (
                    "Excel Macro Enabled",
                    "*.xlsm"
                )
            ]
        )

        if path:

            self.output_file = path

            self.output_var.set(
                path
            )

    # ========================================================
    # LOAD EXCEL
    # ========================================================

    def load_excel(self):

        extension = os.path.splitext(
            self.input_file
        )[1].lower()

        if extension == ".xlsm":

            self.workbook = load_workbook(
                self.input_file,
                keep_vba=True
            )

        else:

            self.workbook = load_workbook(
                self.input_file
            )

    # ========================================================
    # GET ANALYSIS
    # ========================================================

    def get_analysis_data(self):

        try:

            if not self.input_var.get().strip():

                messagebox.showwarning(
                    "Input Required",
                    "Please select an Input Excel File."
                )

                return

            if not self.output_var.get().strip():

                messagebox.showwarning(
                    "Output Required",
                    "Please select a Final Export File Path."
                )

                return

            self.input_file = (
                self.input_var.get().strip()
            )

            self.output_file = (
                self.output_var.get().strip()
            )

            if not os.path.exists(
                self.input_file
            ):

                messagebox.showerror(
                    "File Not Found",
                    "The selected input Excel file does not exist."
                )

                return

            self.load_excel()

            self.build_columns_list()

            self.show_analysis_window()

        except Exception as e:

            messagebox.showerror(
                "Error",
                f"Unable to load Excel file.\n\n{e}"
            )

    # ========================================================
    # BUILD COLUMN LIST
    # ========================================================

    def build_columns_list(self):

        self.columns = []

        for ws in self.workbook.worksheets:

            for col_index in range(
                1,
                ws.max_column + 1
            ):

                header = ws.cell(
                    row=1,
                    column=col_index
                ).value

                if header is None:
                    header = f"Column {col_index}"

                self.columns.append({
                    "sheet": ws.title,
                    "header": str(header),
                    "column_index": col_index
                })

    # ========================================================
    # STEP 2 ANALYSIS
    #
    # Everything is inside ONE window.
    # ========================================================

    def show_analysis_window(self):

        self.analysis_window = tk.Toplevel(
            self.root
        )

        self.analysis_window.title(
            "Step 2 - Analysis"
        )

        self.analysis_window.geometry(
            "1450x850"
        )

        self.analysis_window.minsize(
            1250,
            750
        )

        # ====================================================
        # TOP SUMMARY
        # ====================================================

        summary = tk.Frame(
            self.analysis_window
        )

        summary.pack(
            fill="x",
            padx=20,
            pady=(10, 5)
        )

        total_sheets = len(
            self.workbook.worksheets
        )

        total_columns = sum(
            ws.max_column
            for ws in self.workbook.worksheets
        )

        tk.Label(
            summary,
            text=f"Total Sheet : {total_sheets}",
            font=("Segoe UI", 12, "bold")
        ).pack(
            side="left",
            padx=(10, 40)
        )

        tk.Label(
            summary,
            text=f"Total Column : {total_columns}",
            font=("Segoe UI", 12, "bold")
        ).pack(
            side="left"
        )

        # ====================================================
        # MAIN HORIZONTAL AREA
        # ====================================================

        main = tk.PanedWindow(
            self.analysis_window,
            orient=tk.HORIZONTAL,
            sashrelief="raised",
            sashwidth=6
        )

        main.pack(
            fill="both",
            expand=True,
            padx=15,
            pady=8
        )

        # ====================================================
        # LEFT COLUMN LIST
        # ====================================================

        left = tk.Frame(
            main,
            bd=1,
            relief="solid"
        )

        main.add(
            left,
            minsize=280
        )

        tk.Label(
            left,
            text="All Column Names",
            font=("Segoe UI", 12, "bold")
        ).pack(
            pady=10
        )

        list_frame = tk.Frame(
            left
        )

        list_frame.pack(
            fill="both",
            expand=True,
            padx=10,
            pady=10
        )

        self.analysis_column_list = tk.Listbox(
            list_frame,
            width=38,
            font=("Segoe UI", 10),
            exportselection=False
        )

        list_scroll = ttk.Scrollbar(
            list_frame,
            orient="vertical",
            command=self.analysis_column_list.yview
        )

        self.analysis_column_list.configure(
            yscrollcommand=list_scroll.set
        )

        self.analysis_column_list.pack(
            side="left",
            fill="both",
            expand=True
        )

        list_scroll.pack(
            side="right",
            fill="y"
        )

        for item in self.columns:

            display = (
                f"{item['sheet']} | "
                f"{item['header']}"
            )

            self.analysis_column_list.insert(
                tk.END,
                display
            )

        self.analysis_column_list.bind(
            "<<ListboxSelect>>",
            self.on_analysis_column_selected
        )

        # ====================================================
        # RIGHT ANALYSIS AREA
        # ====================================================

        right = tk.Frame(
            main,
            bd=1,
            relief="solid"
        )

        main.add(
            right,
            minsize=900
        )

        # ====================================================
        # SELECTED COLUMN TITLE
        # ====================================================

        self.analysis_selected_title = tk.Label(
            right,
            text="Column Analysis Details",
            font=("Segoe UI", 13, "bold"),
            anchor="w"
        )

        self.analysis_selected_title.pack(
            fill="x",
            padx=12,
            pady=(10, 5)
        )

        # ====================================================
        # BASIC DETAILS
        # ====================================================

        basic_frame = tk.LabelFrame(
            right,
            text="Column Details",
            font=("Segoe UI", 10, "bold")
        )

        basic_frame.pack(
            fill="x",
            padx=12,
            pady=5
        )

        self.analysis_basic_text = tk.Text(
            basic_frame,
            height=7,
            font=("Consolas", 10),
            wrap="word"
        )

        basic_scroll = ttk.Scrollbar(
            basic_frame,
            orient="vertical",
            command=self.analysis_basic_text.yview
        )

        self.analysis_basic_text.configure(
            yscrollcommand=basic_scroll.set
        )

        self.analysis_basic_text.pack(
            side="left",
            fill="both",
            expand=True,
            padx=(8, 0),
            pady=8
        )

        basic_scroll.pack(
            side="right",
            fill="y",
            pady=8
        )

        self.analysis_basic_text.config(
            state="disabled"
        )

        # ====================================================
        # SPECIAL + NUMERIC AREA
        # ====================================================

        middle = tk.PanedWindow(
            right,
            orient=tk.HORIZONTAL,
            sashrelief="raised",
            sashwidth=5
        )

        middle.pack(
            fill="both",
            expand=True,
            padx=12,
            pady=5
        )

        # ----------------------------------------------------
        # SPECIAL CHARACTER
        # ----------------------------------------------------

        special_frame = tk.LabelFrame(
            middle,
            text="Special Character Present in Column",
            font=("Segoe UI", 10, "bold")
        )

        middle.add(
            special_frame,
            minsize=350
        )

        special_table_frame = tk.Frame(
            special_frame
        )

        special_table_frame.pack(
            fill="both",
            expand=True,
            padx=8,
            pady=8
        )

        self.special_tree = ttk.Treeview(
            special_table_frame,
            columns=(
                "character",
                "occurrence"
            ),
            show="headings"
        )

        self.special_tree.heading(
            "character",
            text="Character"
        )

        self.special_tree.heading(
            "occurrence",
            text="Occurrence Count"
        )

        self.special_tree.column(
            "character",
            width=130,
            anchor="center"
        )

        self.special_tree.column(
            "occurrence",
            width=170,
            anchor="center"
        )

        special_scroll = ttk.Scrollbar(
            special_table_frame,
            orient="vertical",
            command=self.special_tree.yview
        )

        self.special_tree.configure(
            yscrollcommand=special_scroll.set
        )

        self.special_tree.pack(
            side="left",
            fill="both",
            expand=True
        )

        special_scroll.pack(
            side="right",
            fill="y"
        )

        # ----------------------------------------------------
        # NUMERIC VALUES
        # ----------------------------------------------------

        numeric_frame = tk.LabelFrame(
            middle,
            text="Numeric Value in Column",
            font=("Segoe UI", 10, "bold")
        )

        middle.add(
            numeric_frame,
            minsize=350
        )

        numeric_table_frame = tk.Frame(
            numeric_frame
        )

        numeric_table_frame.pack(
            fill="both",
            expand=True,
            padx=8,
            pady=8
        )

        self.numeric_tree = ttk.Treeview(
            numeric_table_frame,
            columns=(
                "number",
                "occurrence"
            ),
            show="headings"
        )

        self.numeric_tree.heading(
            "number",
            text="Numeric Value"
        )

        self.numeric_tree.heading(
            "occurrence",
            text="Occurrence Count"
        )

        self.numeric_tree.column(
            "number",
            width=180,
            anchor="center"
        )

        self.numeric_tree.column(
            "occurrence",
            width=170,
            anchor="center"
        )

        numeric_scroll = ttk.Scrollbar(
            numeric_table_frame,
            orient="vertical",
            command=self.numeric_tree.yview
        )

        self.numeric_tree.configure(
            yscrollcommand=numeric_scroll.set
        )

        self.numeric_tree.pack(
            side="left",
            fill="both",
            expand=True
        )

        numeric_scroll.pack(
            side="right",
            fill="y"
        )

        # ====================================================
        # DOUBLE SPACE
        #
        # This is part of the SAME Analysis window.
        # ====================================================

        double_frame = tk.LabelFrame(
            right,
            text="Double Space in Column",
            font=("Segoe UI", 10, "bold")
        )

        double_frame.pack(
            fill="both",
            expand=True,
            padx=12,
            pady=5
        )

        double_table_frame = tk.Frame(
            double_frame
        )

        double_table_frame.pack(
            fill="both",
            expand=True,
            padx=8,
            pady=8
        )

        self.double_space_tree = ttk.Treeview(
            double_table_frame,
            columns=(
                "row",
                "cell",
                "value"
            ),
            show="headings"
        )

        self.double_space_tree.heading(
            "row",
            text="Row"
        )

        self.double_space_tree.heading(
            "cell",
            text="Cell"
        )

        self.double_space_tree.heading(
            "value",
            text="Value"
        )

        self.double_space_tree.column(
            "row",
            width=100,
            anchor="center"
        )

        self.double_space_tree.column(
            "cell",
            width=120,
            anchor="center"
        )

        self.double_space_tree.column(
            "value",
            width=700,
            anchor="w"
        )

        double_y = ttk.Scrollbar(
            double_table_frame,
            orient="vertical",
            command=self.double_space_tree.yview
        )

        double_x = ttk.Scrollbar(
            double_table_frame,
            orient="horizontal",
            command=self.double_space_tree.xview
        )

        self.double_space_tree.configure(
            yscrollcommand=double_y.set,
            xscrollcommand=double_x.set
        )

        self.double_space_tree.pack(
            side="top",
            fill="both",
            expand=True
        )

        double_y.pack(
            side="right",
            fill="y"
        )

        double_x.pack(
            side="bottom",
            fill="x"
        )

        # ====================================================
        # NEXT BUTTON
        # ====================================================

        ttk.Button(
            self.analysis_window,
            text="Next to Format Data Step",
            command=self.open_format_step
        ).pack(
            pady=(5, 12)
        )

        # ====================================================
        # SELECT FIRST COLUMN
        # ====================================================

        if self.columns:

            self.analysis_column_list.selection_set(
                0
            )

            self.analysis_column_list.event_generate(
                "<<ListboxSelect>>"
            )

    # ========================================================
    # ANALYSIS COLUMN SELECTED
    # ========================================================

    def on_analysis_column_selected(
        self,
        event=None
    ):

        selection = (
            self.analysis_column_list.curselection()
        )

        if not selection:
            return

        index = selection[0]

        column_info = self.columns[index]

        ws = self.workbook[
            column_info["sheet"]
        ]

        col_index = column_info[
            "column_index"
        ]

        value_count = 0
        empty_count = 0

        duplicate_counter = Counter()

        special_counter = Counter()

        numeric_counter = Counter()

        double_spaces = []

        # ====================================================
        # ANALYZE
        # ====================================================

        for row in range(
            2,
            ws.max_row + 1
        ):

            cell = ws.cell(
                row=row,
                column=col_index
            )

            value = cell.value

            if is_blank(value):

                empty_count += 1

                continue

            value_count += 1

            text = value_to_text(
                value
            )

            duplicate_counter[text] += 1

            # ------------------------------------------------
            # SPECIAL CHARACTERS
            # ------------------------------------------------

            for char in text:

                if (
                    not char.isalnum()
                    and not char.isspace()
                ):

                    special_counter[char] += 1

            # ------------------------------------------------
            # NUMERIC VALUES
            # ------------------------------------------------

            if is_numeric_value(value):

                numeric_counter[
                    value_to_text(value)
                ] += 1

            # ------------------------------------------------
            # DOUBLE SPACES
            # ------------------------------------------------

            if isinstance(
                value,
                str
            ):

                if re.search(
                    r" {2,}",
                    value
                ):

                    double_spaces.append(
                        (
                            row,
                            cell.coordinate,
                            value
                        )
                    )

        # ====================================================
        # DUPLICATES
        # ====================================================

        duplicate_value_count = sum(
            count - 1
            for count in duplicate_counter.values()
            if count > 1
        )

        # ====================================================
        # COLUMN TITLE
        # ====================================================

        self.analysis_selected_title.config(
            text=(
                f"Column Analysis: "
                f"{column_info['sheet']} | "
                f"{column_info['header']}"
            )
        )

        # ====================================================
        # BASIC DETAILS
        # ====================================================

        basic_details = (
            f"Column Name              : "
            f"{column_info['header']}\n"

            f"Sheet Name               : "
            f"{column_info['sheet']}\n"

            f"Column Value Count       : "
            f"{value_count}\n"

            f"Empty Cell Count         : "
            f"{empty_count}\n"

            f"Column Duplicate Value Count : "
            f"{duplicate_value_count}\n"

            f"Double Space Cell Count  : "
            f"{len(double_spaces)}"
        )

        self.analysis_basic_text.config(
            state="normal"
        )

        self.analysis_basic_text.delete(
            "1.0",
            tk.END
        )

        self.analysis_basic_text.insert(
            tk.END,
            basic_details
        )

        self.analysis_basic_text.config(
            state="disabled"
        )

        # ====================================================
        # CLEAR SPECIAL TABLE
        # ====================================================

        for item in self.special_tree.get_children():

            self.special_tree.delete(
                item
            )

        # ====================================================
        # SPECIAL CHARACTER DATA
        # ====================================================

        for char, count in sorted(
            special_counter.items(),
            key=lambda x: (
                -x[1],
                x[0]
            )
        ):

            self.special_tree.insert(
                "",
                tk.END,
                values=(
                    repr(char),
                    count
                )
            )

        # ====================================================
        # CLEAR NUMERIC TABLE
        # ====================================================

        for item in self.numeric_tree.get_children():

            self.numeric_tree.delete(
                item
            )

        # ====================================================
        # NUMERIC DATA
        # ====================================================

        for value, count in sorted(
            numeric_counter.items(),
            key=lambda x: (
                -x[1],
                str(x[0])
            )
        ):

            self.numeric_tree.insert(
                "",
                tk.END,
                values=(
                    value,
                    count
                )
            )

        # ====================================================
        # CLEAR DOUBLE SPACE TABLE
        # ====================================================

        for item in self.double_space_tree.get_children():

            self.double_space_tree.delete(
                item
            )

        # ====================================================
        # DOUBLE SPACE DATA
        # ====================================================

        for row, cell_ref, value in double_spaces:

            self.double_space_tree.insert(
                "",
                tk.END,
                values=(
                    row,
                    cell_ref,
                    value
                )
            )

    # ========================================================
    # STEP 3
    # ========================================================

    def open_format_step(self):

        self.analysis_window.destroy()

        self.format_results = []

        self.format_window = tk.Toplevel(
            self.root
        )

        self.format_window.title(
            "Step 3 - Format Data"
        )

        self.format_window.geometry(
            "1400x820"
        )

        self.format_window.minsize(
            1200,
            700
        )

        tk.Label(
            self.format_window,
            text="Step 3 - Format Data",
            font=("Segoe UI", 18, "bold")
        ).pack(
            pady=10
        )

        main = tk.Frame(
            self.format_window
        )

        main.pack(
            fill="both",
            expand=True,
            padx=10,
            pady=5
        )

        # ====================================================
        # LEFT
        # ====================================================

        left = tk.Frame(
            main,
            bd=1,
            relief="solid"
        )

        left.pack(
            side="left",
            fill="y",
            padx=(0, 8)
        )

        tk.Label(
            left,
            text="All Column Names",
            font=("Segoe UI", 11, "bold")
        ).pack(
            pady=8
        )

        column_frame = tk.Frame(
            left
        )

        column_frame.pack(
            fill="both",
            expand=True,
            padx=8,
            pady=8
        )

        self.format_column_list = tk.Listbox(
            column_frame,
            width=32,
            exportselection=False
        )

        format_list_scroll = ttk.Scrollbar(
            column_frame,
            orient="vertical",
            command=self.format_column_list.yview
        )

        self.format_column_list.configure(
            yscrollcommand=format_list_scroll.set
        )

        self.format_column_list.pack(
            side="left",
            fill="both",
            expand=True
        )

        format_list_scroll.pack(
            side="right",
            fill="y"
        )

        for item in self.columns:

            display = (
                f"{item['sheet']} | "
                f"{item['header']}"
            )

            self.format_column_list.insert(
                tk.END,
                display
            )

        self.format_column_list.bind(
            "<<ListboxSelect>>",
            self.on_format_column_selected
        )

        # ====================================================
        # MIDDLE
        # ====================================================

        middle = tk.Frame(
            main,
            bd=1,
            relief="solid"
        )

        middle.pack(
            side="left",
            fill="both",
            expand=True,
            padx=(0, 8)
        )

        tk.Label(
            middle,
            text="Full Selected Column Data",
            font=("Segoe UI", 11, "bold")
        ).pack(
            pady=8
        )

        data_frame = tk.Frame(
            middle
        )

        data_frame.pack(
            fill="both",
            expand=True,
            padx=8,
            pady=8
        )

        self.column_data_tree = ttk.Treeview(
            data_frame,
            columns=(
                "row",
                "cell",
                "value"
            ),
            show="headings"
        )

        self.column_data_tree.heading(
            "row",
            text="Row"
        )

        self.column_data_tree.heading(
            "cell",
            text="Cell"
        )

        self.column_data_tree.heading(
            "value",
            text="Value"
        )

        self.column_data_tree.column(
            "row",
            width=70
        )

        self.column_data_tree.column(
            "cell",
            width=90
        )

        self.column_data_tree.column(
            "value",
            width=420
        )

        data_scroll = ttk.Scrollbar(
            data_frame,
            orient="vertical",
            command=self.column_data_tree.yview
        )

        self.column_data_tree.configure(
            yscrollcommand=data_scroll.set
        )

        self.column_data_tree.pack(
            side="left",
            fill="both",
            expand=True
        )

        data_scroll.pack(
            side="right",
            fill="y"
        )

        # ====================================================
        # RIGHT
        # ====================================================

        right = tk.Frame(
            main,
            bd=1,
            relief="solid"
        )

        right.pack(
            side="left",
            fill="both",
            expand=True
        )

        tk.Label(
            right,
            text="Select One Task",
            font=("Segoe UI", 11, "bold")
        ).pack(
            pady=8
        )

        self.task_var = tk.StringVar()

        self.task_combo = ttk.Combobox(
            right,
            textvariable=self.task_var,
            values=TASKS,
            state="readonly",
            width=38
        )

        self.task_combo.pack(
            padx=10,
            pady=5
        )

        ttk.Button(
            right,
            text="Apply",
            command=self.apply_selected_task
        ).pack(
            pady=8
        )

        tk.Label(
            right,
            text="Task Performed",
            font=("Segoe UI", 11, "bold")
        ).pack(
            pady=(12, 5)
        )

        info_frame = tk.Frame(
            right
        )

        info_frame.pack(
            fill="x",
            padx=10
        )

        self.task_info_text = tk.Text(
            info_frame,
            height=6,
            wrap="word"
        )

        self.task_info_text.pack(
            fill="x"
        )

        self.task_info_text.config(
            state="disabled"
        )

        tk.Label(
            right,
            text="Changed Value Result",
            font=("Segoe UI", 11, "bold")
        ).pack(
            pady=(12, 5)
        )

        result_frame = tk.Frame(
            right
        )

        result_frame.pack(
            fill="both",
            expand=True,
            padx=10,
            pady=(0, 10)
        )

        self.result_tree = ttk.Treeview(
            result_frame,
            columns=(
                "old",
                "new",
                "task"
            ),
            show="headings"
        )

        self.result_tree.heading(
            "old",
            text="Old Value"
        )

        self.result_tree.heading(
            "new",
            text="New Value"
        )

        self.result_tree.heading(
            "task",
            text="Task Name"
        )

        self.result_tree.column(
            "old",
            width=170
        )

        self.result_tree.column(
            "new",
            width=170
        )

        self.result_tree.column(
            "task",
            width=190
        )

        result_scroll = ttk.Scrollbar(
            result_frame,
            orient="vertical",
            command=self.result_tree.yview
        )

        self.result_tree.configure(
            yscrollcommand=result_scroll.set
        )

        self.result_tree.pack(
            side="left",
            fill="both",
            expand=True
        )

        result_scroll.pack(
            side="right",
            fill="y"
        )

        # ====================================================
        # BOTTOM
        # ====================================================

        bottom = tk.Frame(
            self.format_window
        )

        bottom.pack(
            fill="x",
            pady=10
        )

        ttk.Button(
            bottom,
            text="Final Export",
            command=self.final_export
        ).pack(
            side="right",
            padx=20
        )

        ttk.Button(
            bottom,
            text="Back to Analysis",
            command=self.back_to_analysis
        ).pack(
            side="right",
            padx=10
        )

        if self.columns:

            self.format_column_list.selection_set(
                0
            )

            self.format_column_list.event_generate(
                "<<ListboxSelect>>"
            )

    # ========================================================
    # FORMAT COLUMN SELECTED
    # ========================================================

    def on_format_column_selected(
        self,
        event=None
    ):

        selection = (
            self.format_column_list.curselection()
        )

        if not selection:
            return

        index = selection[0]

        self.current_column = self.columns[
            index
        ]

        self.show_selected_column_data()

        self.update_task_info("")

    # ========================================================
    # SHOW COLUMN DATA
    # ========================================================

    def show_selected_column_data(self):

        for item in self.column_data_tree.get_children():

            self.column_data_tree.delete(
                item
            )

        if not self.current_column:
            return

        ws = self.workbook[
            self.current_column["sheet"]
        ]

        col_index = self.current_column[
            "column_index"
        ]

        for row in range(
            1,
            ws.max_row + 1
        ):

            cell = ws.cell(
                row=row,
                column=col_index
            )

            self.column_data_tree.insert(
                "",
                tk.END,
                values=(
                    row,
                    cell.coordinate,
                    value_to_text(
                        cell.value
                    )
                )
            )

    # ========================================================
    # TASK INFO
    # ========================================================

    def update_task_info(self, text):

        self.task_info_text.config(
            state="normal"
        )

        self.task_info_text.delete(
            "1.0",
            tk.END
        )

        self.task_info_text.insert(
            tk.END,
            text
        )

        self.task_info_text.config(
            state="disabled"
        )

    # ========================================================
    # APPLY TASK
    # ========================================================

    def apply_selected_task(self):

        if not self.current_column:

            messagebox.showwarning(
                "Column Required",
                "Please select a column first."
            )

            return

        task = self.task_var.get()

        if not task:

            messagebox.showwarning(
                "Task Required",
                "Please select a task."
            )

            return

        if task == "Length Count":

            self.length_count_task()

            return

        try:

            ws = self.workbook[
                self.current_column["sheet"]
            ]

            col_index = self.current_column[
                "column_index"
            ]

            changed_count = 0

            task_changes = []

            for row in range(
                2,
                ws.max_row + 1
            ):

                cell = ws.cell(
                    row=row,
                    column=col_index
                )

                old_value = cell.value

                # ------------------------------------------------
                # TRIM
                # ------------------------------------------------

                if task == "Trim":

                    new_value = trim_excel_style(
                        old_value
                    )

                # ------------------------------------------------
                # PROPER CASE
                # ------------------------------------------------

                elif task == "Poper Case":

                    new_value = proper_case(
                        old_value
                    )

                # ------------------------------------------------
                # UPPER CASE
                # ------------------------------------------------

                elif task == "Upper Case":

                    new_value = upper_case(
                        old_value
                    )

                # ------------------------------------------------
                # REMOVE SPECIAL
                # ------------------------------------------------

                elif task == "Remove Special Character":

                    new_value = remove_selected_specials(
                        old_value
                    )

                # ------------------------------------------------
                # SSN
                # ------------------------------------------------

                elif task == "SSN Format":

                    new_value = format_ssn(
                        old_value
                    )

                # ------------------------------------------------
                # ITIN
                # ------------------------------------------------

                elif task == "ITIN Find":

                    if starts_with_nine(
                        old_value
                    ):

                        cell.fill = BLUE_FILL

                    continue

                # ------------------------------------------------
                # INVALID SSN
                # ------------------------------------------------

                elif task == "Find Invalid SSN":

                    if invalid_ssn(
                        old_value
                    ):

                        cell.fill = PINK_FILL

                    continue

                # ------------------------------------------------
                # DOB
                # ------------------------------------------------

                elif task == "DOB (mm/dd/yyyy)":

                    new_value = format_dob(
                        old_value
                    )

                # ------------------------------------------------
                # PHONE
                # ------------------------------------------------

                elif task == "Phone Number (123-456-7890)":

                    new_value = format_phone(
                        old_value
                    )

                # ------------------------------------------------
                # STREET
                # ------------------------------------------------

                elif task == "Street Address Format":

                    new_value = street_address_format(
                        old_value
                    )

                else:

                    new_value = old_value

                # ------------------------------------------------
                # ONLY CHANGE ACTUAL CHANGES
                # ------------------------------------------------

                if new_value != old_value:

                    cell.value = new_value

                    cell.fill = YELLOW_FILL

                    changed_count += 1

                    task_changes.append({
                        "old": value_to_text(
                            old_value
                        ),
                        "new": value_to_text(
                            new_value
                        ),
                        "task": task
                    })

            self.format_results.extend(
                task_changes
            )

            for change in task_changes:

                self.result_tree.insert(
                    "",
                    tk.END,
                    values=(
                        change["old"],
                        change["new"],
                        change["task"]
                    )
                )

            self.update_task_info(
                f"Task Performed Name: {task}\n"
                f"Column Name: "
                f"{self.current_column['header']}\n"
                f"Sheet Name: "
                f"{self.current_column['sheet']}\n"
                f"Changed Value Count: {changed_count}"
            )

            self.show_selected_column_data()

            if task == "ITIN Find":

                messagebox.showinfo(
                    "ITIN Find",
                    "ITIN check completed.\n\n"
                    "Matching cells were highlighted BLUE."
                )

            elif task == "Find Invalid SSN":

                messagebox.showinfo(
                    "Invalid SSN",
                    "Invalid SSN check completed.\n\n"
                    "Matching cells were highlighted PINK."
                )

            else:

                messagebox.showinfo(
                    "Task Completed",
                    f"{task}\n\n"
                    f"Column: "
                    f"{self.current_column['header']}\n"
                    f"Changed Value Count: "
                    f"{changed_count}"
                )

        except Exception as e:

            messagebox.showerror(
                "Formatting Error",
                f"Unable to perform task.\n\n{e}"
            )

    # ========================================================
    # LENGTH COUNT
    # ========================================================

    def length_count_task(self):

        if not self.current_column:
            return

        ws = self.workbook[
            self.current_column["sheet"]
        ]

        col_index = self.current_column[
            "column_index"
        ]

        length_counter = Counter()

        for row in range(
            2,
            ws.max_row + 1
        ):

            value = ws.cell(
                row=row,
                column=col_index
            ).value

            if is_blank(value):
                continue

            text = value_to_text(
                value
            )

            length_counter[
                len(text)
            ] += 1

        popup = tk.Toplevel(
            self.format_window
        )

        popup.title(
            "Length Count"
        )

        popup.geometry(
            "500x500"
        )

        tk.Label(
            popup,
            text="Length Count Summary",
            font=("Segoe UI", 14, "bold")
        ).pack(
            pady=15
        )

        tk.Label(
            popup,
            text=(
                f"Column: "
                f"{self.current_column['header']}"
            ),
            font=("Segoe UI", 10)
        ).pack(
            pady=(0, 10)
        )

        frame = tk.Frame(
            popup
        )

        frame.pack(
            fill="both",
            expand=True,
            padx=15,
            pady=10
        )

        tree = ttk.Treeview(
            frame,
            columns=(
                "count",
                "length"
            ),
            show="headings"
        )

        tree.heading(
            "count",
            text="Value Count"
        )

        tree.heading(
            "length",
            text="Length"
        )

        tree.column(
            "count",
            width=180
        )

        tree.column(
            "length",
            width=180
        )

        scrollbar = ttk.Scrollbar(
            frame,
            orient="vertical",
            command=tree.yview
        )

        tree.configure(
            yscrollcommand=scrollbar.set
        )

        tree.pack(
            side="left",
            fill="both",
            expand=True
        )

        scrollbar.pack(
            side="right",
            fill="y"
        )

        sorted_lengths = sorted(
            length_counter.items(),
            key=lambda x: (
                -x[1],
                -x[0]
            )
        )

        for length, count in sorted_lengths:

            tree.insert(
                "",
                tk.END,
                values=(
                    count,
                    length
                )
            )

        ttk.Button(
            popup,
            text="Close",
            command=popup.destroy
        ).pack(
            pady=15
        )

    # ========================================================
    # BACK
    # ========================================================

    def back_to_analysis(self):

        if self.format_window:

            self.format_window.destroy()

        self.show_analysis_window()

    # ========================================================
    # FINAL EXPORT
    # ========================================================

    def final_export(self):

        try:

            output = self.output_var.get().strip()

            if not output:

                messagebox.showwarning(
                    "Output Required",
                    "Please select a Final Export File Path."
                )

                return

            if (
                os.path.abspath(output)
                ==
                os.path.abspath(
                    self.input_file
                )
            ):

                messagebox.showerror(
                    "Invalid Export Path",
                    "The Final Export path cannot be "
                    "the same as the original input file."
                )

                return

            directory = os.path.dirname(
                os.path.abspath(output)
            )

            if not os.path.exists(
                directory
            ):

                os.makedirs(
                    directory,
                    exist_ok=True
                )

            self.workbook.save(
                output
            )

            total_changes = len(
                self.format_results
            )

            messagebox.showinfo(
                "Export Successful",
                "Excel file exported successfully.\n\n"
                f"Location:\n{output}\n\n"
                f"Total Changed Values: "
                f"{total_changes}\n\n"
                "Original input file was not modified."
            )

        except PermissionError:

            messagebox.showerror(
                "Permission Error",
                "The output Excel file may already be open.\n\n"
                "Please close it and try again."
            )

        except Exception as e:

            messagebox.showerror(
                "Export Error",
                f"Unable to export Excel file.\n\n{e}"
            )


# ============================================================
# MAIN
# ============================================================

def main():

    root = tk.Tk()

    try:

        style = ttk.Style()

        if "vista" in style.theme_names():

            style.theme_use("vista")

        elif "clam" in style.theme_names():

            style.theme_use("clam")

    except Exception:
        pass

    ExcelFormatApp(root)

    root.mainloop()


if __name__ == "__main__":
    main()